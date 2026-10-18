"""Local document parsing: PDF (PyMuPDF), DOCX (python-docx), TXT.

All parsing is in-memory from bytes. Errors are raised as IngestError with a user-facing
message, never a stack trace.
"""
from __future__ import annotations

import io
import logging
import re
import unicodedata

from lexlocal import config
from lexlocal.ingest import adversarial
from lexlocal.models import AdversarialFlag, Document, Page

log = logging.getLogger("lexlocal.ingest")


class IngestError(Exception):
    """User-facing ingestion failure. `code` is machine readable for tests/eval."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code
        self.message = message


CONTRACT_MARKERS = [
    "agreement", "lease", "tenant", "landlord", "lessor", "lessee", "licensor", "licensee",
    "employer", "employee", "employment", "party", "parties", "hereby", "whereas", "clause",
    "terminate", "termination", "notice", "rent", "salary", "deposit", "witness", "executed",
    "obligation", "shall", "term of", "consideration", "agree",
]



def normalize_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = adversarial.INVISIBLE_RE.sub("", text)  # zero-width and bidi controls
    text = adversarial.normalize_homoglyphs(text)
    # Some fonts extract "-" as a soft hyphen (U+00AD): drop it at line ends, keep it as "-" elsewhere.
    text = text.replace(chr(0xAD) + chr(10), "").replace(chr(0xAD), "-")
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace(" ", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"-\n(?=[a-z])", "", text)  # de-hyphenate line breaks
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _assemble(name: str, kind: str, page_texts: list[str], warnings: list[str],
              flags: list[AdversarialFlag] | None = None) -> Document:
    pages: list[Page] = []
    parts: list[str] = []
    flags = list(flags or [])
    offset = 0
    for i, raw in enumerate(page_texts, start=1):
        flags += adversarial.scan_text(raw, i)
        t = normalize_text(raw)
        pages.append(Page(number=i, text=t, start=offset, end=offset + len(t) + 2))
        parts.append(t)
        offset += len(t) + 2
    text = "\n\n".join(parts)
    if flags:
        warnings = warnings + [f"Content a reader would not see was found ({len(flags)} item(s)) and kept out of the "
                               "analysis. See Safety and privacy."]
    return Document(name=name, kind=kind, pages=pages, text=text, warnings=warnings, adversarial=flags)


def _parse_pdf(name: str, data: bytes, vault=None) -> Document:
    import pymupdf as fitz

    try:
        doc = fitz.open(stream=data, filetype="pdf")
    except Exception:
        raise IngestError("corrupt", "This PDF appears to be damaged or is not a valid PDF, so it could not be opened.")
    try:
        if doc.needs_pass or doc.is_encrypted:
            # Try empty owner password (common for 'print-protected' PDFs)
            if not doc.authenticate(""):
                raise IngestError("encrypted", "This PDF is password-protected. Please remove the password locally and upload it again.")
        if doc.page_count == 0:
            raise IngestError("empty", "This PDF has no pages.")
        if doc.page_count > config.MAX_PAGES:
            raise IngestError("too_many_pages", f"This PDF has {doc.page_count} pages; the limit is {config.MAX_PAGES}.")
        page_texts, image_pages, flags = [], [], []
        for page in doc:
            t = page.get_text("text", sort=True) or ""
            page_flags, hidden = adversarial.scan_pdf_page(page)
            flags += page_flags
            t = adversarial.remove_hidden(t, hidden)
            page_texts.append(t)
            if len(t.strip()) < config.MIN_TEXT_CHARS_PER_PAGE and page.get_images(full=False):
                image_pages.append(page.number)
        ocr_conf, ocr_note = _ocr_pages(doc, image_pages, page_texts, vault)
    except IngestError:
        raise
    except Exception:
        raise IngestError("corrupt", "This PDF could not be read; it may be damaged.")
    finally:
        doc.close()

    total = sum(len(t.strip()) for t in page_texts)
    low_text_pages = sum(1 for t in page_texts if len(t.strip()) < config.MIN_TEXT_CHARS_PER_PAGE)
    if total < config.MIN_TEXT_CHARS_PER_PAGE:
        if image_pages:
            raise IngestError("scanned", "This looks like a scanned PDF (page images, no text layer). " + ocr_note)
        raise IngestError("empty", "No readable text was found in this PDF.")
    warnings = []
    if ocr_conf:
        low = sorted(p for p, c in ocr_conf.items() if c < config.OCR_LOW_CONFIDENCE)
        mean = sum(ocr_conf.values()) / len(ocr_conf)
        warnings.append(f"{len(ocr_conf)} scanned page(s) were read with local OCR (average confidence {mean:.0%}). "
                        "Quotes from these pages come from OCR text, so check them against the page image.")
        if low:
            warnings.append(f"OCR confidence is low on page(s) {', '.join(map(str, low))}. Findings from those pages "
                            "have lower confidence; check them against the original.")
    elif ocr_note and image_pages:
        warnings.append(ocr_note)
    if low_text_pages and low_text_pages >= max(1, len(page_texts) // 2):
        warnings.append(f"{low_text_pages} of {len(page_texts)} pages have little or no text. Those pages were not analysed.")
    d = _assemble(name, "pdf", page_texts, warnings, flags)
    d.ocr_confidence = ocr_conf
    return d


def _ocr_pages(doc, page_numbers: list[int], page_texts: list[str], vault) -> tuple[dict[int, float], str]:
    """OCR image-only pages in place. Returns ({1-based page: confidence}, note for the user)."""
    if not page_numbers:
        return {}, ""
    from lexlocal.ingest import ocr

    if not config.OCR_ENABLED:
        return {}, "Local OCR is turned off (LEXLOCAL_OCR=off), so scanned pages were not read."
    ok, msg = ocr.available()
    if not ok:
        return {}, msg
    own_vault = vault is None
    if own_vault:
        from lexlocal.privacy.vault import Vault

        vault = Vault()
    conf = {}
    try:
        for n in page_numbers[:config.MAX_OCR_PAGES]:
            r = ocr.ocr_page(doc[n], vault)
            page_texts[n] = r.text
            conf[n + 1] = r.confidence
    finally:
        if own_vault:
            vault.wipe()
    skipped = len(page_numbers) - len(conf)
    note = f" {skipped} further scanned page(s) were not read (limit {config.MAX_OCR_PAGES})." if skipped else ""
    return conf, note.strip()


def _parse_docx(name: str, data: bytes) -> Document:
    try:
        import docx

        d = docx.Document(io.BytesIO(data))
    except Exception:
        raise IngestError("corrupt", "This DOCX file could not be opened; it may be damaged or not a Word document.")
    lines = [p.text for p in d.paragraphs]
    for table in d.tables:
        for row in table.rows:
            lines.append(" | ".join(c.text.strip() for c in row.cells))
    text = "\n".join(lines)
    # DOCX has no fixed pages: split into pseudo-pages of ~3000 chars on paragraph boundaries
    pages, cur = [], ""
    for para in text.split("\n"):
        if len(cur) + len(para) > 3000 and cur:
            pages.append(cur)
            cur = ""
        cur += para + "\n"
    if cur.strip():
        pages.append(cur)
    if not "".join(pages).strip():
        raise IngestError("empty", "This DOCX file contains no text.")
    return _assemble(name, "docx", pages, ["DOCX has no fixed pages; page numbers refer to ~3000-character sections."])


def _parse_txt(name: str, data: bytes) -> Document:
    if b"\x00" in data[:4096]:
        raise IngestError("corrupt", "This file looks binary, not plain text.")
    for enc in ("utf-8-sig", "utf-16", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if not text.strip():
        raise IngestError("empty", "This text file is empty.")
    # Form feed = page break; otherwise pseudo-pages of ~3000 chars
    if "\f" in text:
        pages = text.split("\f")
    else:
        pages, cur = [], ""
        for para in text.split("\n"):
            if len(cur) + len(para) > 3000 and cur:
                pages.append(cur)
                cur = ""
            cur += para + "\n"
        pages.append(cur)
    return _assemble(name, "txt", [p for p in pages if p.strip()] or [text], [])


def looks_like_contract(text: str) -> tuple[bool, int]:
    low = text.lower()
    hits = sum(1 for m in CONTRACT_MARKERS if m in low)
    return hits >= 5, hits


def parse_bytes(name: str, data: bytes, vault=None) -> Document:
    if data is None or len(data) == 0:
        raise IngestError("empty", "The uploaded file is empty (0 bytes).")
    if len(data) > config.MAX_FILE_MB * 1024 * 1024:
        raise IngestError("too_large", f"The file is larger than {config.MAX_FILE_MB} MB.")
    ext = name.lower().rsplit(".", 1)[-1] if "." in name else ""
    if data[:5] == b"%PDF-" or ext == "pdf":
        if data[:5] != b"%PDF-":
            raise IngestError("corrupt", "This file has a .pdf extension but is not a valid PDF.")
        doc = _parse_pdf(name, data, vault)
    elif ext == "docx" or data[:2] == b"PK":
        doc = _parse_docx(name, data)
    elif ext in ("txt", "md", "text", ""):
        doc = _parse_txt(name, data)
    else:
        raise IngestError("unsupported", f"Unsupported file type '.{ext}'. Please upload PDF, DOCX or TXT.")

    is_contract, hits = looks_like_contract(doc.text)
    if not is_contract:
        doc.warnings.append(
            "This does not look like a contract or agreement (few legal terms found). "
            "LexLocal is designed for rental/lease and employment agreements; results may not be meaningful."
        )
    log.info("parsed doc kind=%s pages=%d chars=%d contract_markers=%d", doc.kind, len(doc.pages), len(doc.text), hits)
    return doc


def parse_file(path) -> Document:
    from pathlib import Path

    p = Path(path)
    return parse_bytes(p.name, p.read_bytes())
